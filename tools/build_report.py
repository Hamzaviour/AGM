"""
Build the assignment report (report/Smart_Food_Delivery_Agent_Report.docx and .pdf)
from the code, the results and the diagrams.

    pip install python-docx Pillow
    python tools/build_report.py

The PDF is made with LibreOffice (soffice) if it is installed, otherwise with
Microsoft Word through docx2pdf (pip install docx2pdf).
"""

import inspect
import json
import os
import shutil
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Fill these in, then run the script again (or edit the Word file directly).
TITLE_INFO = {
    "course": "Agent Based Modeling",
    "assignment": "Assignment 1 (group)",
    "section": "____________",
    "instructor": "____________",
    "date": "October 2026",
}
MEMBERS = [("Ch Hamza Younas", "L1F23BSAI0089"), ("Rida", "L1F23BSAI0091"), ("Muhammad Abdullah", "L1F23BSAI0063")]

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
sys.path.insert(0, os.path.join(ROOT, "code"))
import agents as agents_module  # noqa: E402
from environment import (DEADLINE_MIN, ROUTES, ROUTE_IDS, SCENARIOS,  # noqa: E402
                         base_minutes, make_percept, route_km)

PINK = RGBColor(0xE2, 0x1B, 0x70)
DIAG = os.path.join(ROOT, "diagrams")
OUT_DIR = os.path.join(ROOT, "report")
NAME = "Smart_Food_Delivery_Agent_Report"


def shade(cell, hex_color):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tc.append(shd)


def table(doc, header, rows, widths=None, highlight=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "E21B70")
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            if highlight is not None and k == highlight:
                shade(cells[i], "FDE6EF")
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
    if widths:
        t.autofit = False
        grid = t._tbl.tblGrid
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
            grid.gridCol_lst[i].w = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def figure(doc, name, caption, width=16):
    path = os.path.join(DIAG, name)
    if not os.path.exists(path):
        doc.add_paragraph("[missing figure: %s]" % name)
        return
    doc.add_picture(path, width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def code(doc, text, size=8.5):
    """Monospaced block on a light grey background."""
    p = doc.add_paragraph()
    r = p.add_run(text.rstrip("\n"))
    r.font.name = "Consolas"; r.font.size = Pt(size)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    pf = p.paragraph_format
    pf.left_indent = Cm(0.3); pf.space_before = Pt(2); pf.space_after = Pt(8)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F3F2EE")
    p._p.get_or_add_pPr().append(shd)


def title_page(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("University of Central Punjab"); r.font.size = Pt(16); r.bold = True
    logo = os.path.join(OUT_DIR, "assets", "ucp_logo.png")
    if os.path.exists(logo):
        doc.add_picture(logo, width=Cm(4.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("%s, %s" % (TITLE_INFO["course"], TITLE_INFO["assignment"])).font.size = Pt(13)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Smart Food Delivery Agent"); r.font.size = Pt(28); r.bold = True; r.font.color.rgb = PINK
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Seven agent architectures for a rider delivering from\nCheezious, Johar Town to the University of Central Punjab")
    r.font.size = Pt(13)
    doc.add_paragraph()
    t = doc.add_table(rows=1, cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.rows[0].cells[0].text = "Member"; t.rows[0].cells[1].text = "Registration No."
    for c in t.rows[0].cells:
        shade(c, "E21B70"); c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    for n, reg in MEMBERS:
        cells = t.add_row().cells; cells[0].text = n; cells[1].text = reg
    doc.add_paragraph()
    for k in ("section", "instructor", "date"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("%s: " % k.capitalize()).bold = True
        p.add_run(TITLE_INFO[k])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build():
    with open(os.path.join(ROOT, "results", "results.json"), encoding="utf-8") as f:
        res = json.load(f)
    agents = {a["key"]: a for a in res["agents"]}

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for s in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[s].font.color.rgb = PINK
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)   # A4
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    doc.core_properties.author = ", ".join(n for n, _ in MEMBERS)
    doc.core_properties.title = "Smart Food Delivery Agent"
    title_page(doc)

    # 1 -------------------------------------------------------------------
    doc.add_heading("1. Introduction", 1)
    doc.add_paragraph(
        "Food delivery riders in Lahore make the same decision many times a day: which road to take from the "
        "restaurant to the customer. The fastest road depends on the time of day, the weather and things the "
        "rider cannot see from the restaurant door. This report models one such rider as an intelligent agent. "
        "An order is picked up at Cheezious, Johar Town and must be handed over at the main gate of the "
        "University of Central Punjab (UCP), Khayaban-e-Jinnah, within %d minutes." % DEADLINE_MIN)
    doc.add_paragraph(
        "The same delivery problem is solved by seven single-agent architectures: Simple Reflex, Model-Based "
        "Reflex, Goal-Based, Utility-Based, Learning, Hierarchical and Contextual. Every agent receives exactly "
        "the same percept and is tested in the same three situations. A Python program produces the results, "
        "and an interactive 3D simulation shows each delivery as a rider riding through a model of Johar Town.")

    # 2 -------------------------------------------------------------------
    doc.add_heading("2. The Environment", 1)
    doc.add_heading("2.1 Road map and routes", 2)
    doc.add_paragraph("There are three ways from Cheezious to UCP. Destinations A, B and C are the three routes:")
    table(doc, ["Route", "Road", "Distance", "Free-flow time", "Main risk"],
          [[r, ROUTES[r]["name"], "%.2f km" % route_km(r), "%.1f min" % base_minutes(r),
            {"A": "Jams at the Jinnah signal chowk in rush hour",
             "B": "The Canal underpass floods after about 30 minutes of rain",
             "C": "Longest route, but rarely blocked"}[r]] for r in ROUTE_IDS],
          widths=[1.4, 4.5, 2.0, 2.4, 6.3])
    figure(doc, "0_route_map.png", "Figure 1. Schematic road map (not to scale).")
    doc.add_paragraph(
        "The road map is a graph with 7 nodes (Cheezious, Jinnah signal chowk, Canal Bank Road entry, Canal "
        "underpass, Block G inner street, Maulana Shaukat Ali Road, UCP gate) and 8 road sections, each with a "
        "length in kilometres. Riding time is distance divided by the route's average speed. A jam adds a fixed "
        "delay (9.2 min on A), and a flooded underpass adds 14.1 min to B, where the rider has to push the bike "
        "through the water.")

    doc.add_heading("2.2 PEAS description", 2)
    table(doc, ["", "Description"], [
        ["Performance", "Order handed over within 20 minutes; shortest delivery time; avoid unsafe (flooded) roads; arrive at all (do not choose a closed road)."],
        ["Environment", "Johar Town roads between Cheezious and UCP: routes A, B, C, traffic, rain, road closures, underpass flooding, time of day."],
        ["Actuators", "Motorbike controls: choose and ride route A, B or C; speed and braking (used by the Hierarchical agent's operational layer)."],
        ["Sensors", "Eyes (closures and traffic right outside the restaurant, weather), the rider app (clock, live traffic on all routes, reported closures), memory of earlier percepts."],
    ], widths=[3, 13.6])

    doc.add_heading("2.3 Properties of the task environment", 2)
    table(doc, ["Property", "Value", "Reason"], [
        ["Observable", "Partially", "Flooding of the underpass is 2 km away and is not reported by the app."],
        ["Agents", "Single agent", "Other traffic is part of the environment, not competing agents."],
        ["Deterministic", "Stochastic (to the agent)", "Whether the underpass is flooded is uncertain given the percept; the Learning agent's training trips have random noise."],
        ["Episodic", "Episodic", "Each delivery is independent; the Learning agent uses earlier episodes only for training."],
        ["Static", "Static during the decision", "The world does not change while the agent decides; the route is chosen once at the restaurant."],
        ["Discrete", "Discrete", "Three actions (A, B, C) and a finite set of road conditions."],
        ["Known", "Known", "The rider knows the road map and the rules of the environment."],
    ], widths=[3, 4, 9.6])

    doc.add_heading("2.4 Percept", 2)
    doc.add_paragraph("Every agent receives the same percept. Example from Test 3:")
    p3 = make_percept(SCENARIOS[2])
    code(doc, "\n".join("%-9s %s" % (k + ":", json.dumps(v)) for k, v in p3.items()))

    # 3 -------------------------------------------------------------------
    doc.add_heading("3. Test situations", 1)
    table(doc, ["Test", "Situation", "Order time", "What is true", "What the rider can see"], [
        ["1", "Normal lunch hour", "13:00", "All roads clear", "Everything"],
        ["2", "Evening rush hour", "18:10", "A jammed at Jinnah chowk (+9.2 min)", "A heavy on the traffic app"],
        ["3", "Monsoon rain", "13:00", "Rain since 12:15; A closed by a fallen tree; B underpass flooded (+14.1 min)", "Rain, A closed. Flooding is NOT visible"],
    ], widths=[1.2, 3.2, 2.0, 5.4, 4.8])
    doc.add_paragraph("The best possible choice is A (10.0 min) in Test 1, B (10.4 min) in Test 2 and C (14.4 min) in Test 3.")

    # 4 -------------------------------------------------------------------
    doc.add_heading("4. Agent designs", 1)
    designs = [
        ("simple_reflex", "4.1 Simple Reflex Agent", "1_simple_reflex.png",
         "Uses condition-action rules on the current percept only. It has no memory and no model of the world.",
         "IF A is closed THEN take B\nIF A is jammed THEN take B\nOTHERWISE take A",
         "Fast and simple. In Test 3 it cannot know that the underpass is flooded, so it sends the rider into the water."),
        ("model_based", "4.2 Model-Based Reflex Agent", "2_model_based_reflex.png",
         "Keeps an internal state that it updates from the percept history, using a model of how the world works: "
         "the Canal underpass floods after 30 minutes of continuous rain.",
         "state.rain_minutes = time since rain started (from percept history)\nstate.B_flooded = rain_minutes >= 30\n"
         "IF A not blocked THEN A ELSE IF B not flooded THEN B ELSE C",
         "Handles the partially observable flood by remembering how long it has been raining."),
        ("goal_based", "4.3 Goal-Based Agent", "3_goal_based.png",
         "Has an explicit goal (be at the UCP gate) and searches the road graph with uniform-cost search, using "
         "distance as the cost. Closed roads are removed from the graph.",
         "graph = road map without closed roads\npath = uniform_cost_search(graph, CHZ, UCP)\nride the route of path",
         "Always reaches the goal by the shortest road, but the goal test does not care about traffic or weather. "
         "This costs time in Test 2 and makes it late in Test 3."),
        ("utility_based", "4.4 Utility-Based Agent", "4_utility_based.png",
         "Predicts the possible outcomes of each route with probabilities, for example P(flood) = 0.85 after more "
         "than 30 minutes of rain, and scores them with a utility function:",
         "U(route) = 100 x P(on time) - 2 x E[minutes] - 15 x P(flood)\nchoose the open route with the highest U",
         "Trades off speed, punctuality and safety, and handles uncertainty explicitly."),
        ("learning", "4.5 Learning Agent", "5_learning.png",
         "Learns a value Q(context, route) from experience. During training it rides every route on six practice "
         "days three times (54 trips, with a little random noise) and keeps the running average of -minutes. "
         "Contexts are normal, rush and rain.",
         "Q(ctx, r) <- Q(ctx, r) + (reward - Q(ctx, r)) / n      reward = -minutes\nat test time: choose argmax_r Q(ctx, r) among open routes",
         "Needs no hand-written rules, but it needs training data that covers each context."),
        ("hierarchical", "4.6 Hierarchical Agent", "6_hierarchical.png",
         "Splits the problem into three layers. The strategic layer sets the mode (fast or safe) and constraints "
         "(avoid congestion, avoid flood-prone roads after 30 minutes of rain). The tactical layer picks the fastest "
         "route that meets the constraints. The operational layer controls riding (speed limit, signals).",
         "constraints = strategic(percept)\nroute = fastest route that satisfies constraints\noperational(route): ride it",
         "Easy to extend: a new concern (for example fuel) is added in one layer."),
        ("contextual", "4.7 Contextual Agent", "7_contextual.png",
         "First recognises the context (normal, rush hour or monsoon) from the clock, weather and traffic, then "
         "applies the policy written for that context.",
         "normal: A > B > C     rush hour: B > C > A     monsoon: C > A > B\nchoose the first open route in the policy",
         "Very quick decisions, but every context and its policy has to be designed in advance."),
    ]
    classes = {c.key: c for c in agents_module.ALL_AGENTS}
    for key, title, img, desc, rules, note in designs:
        cls = classes[key]
        name = title.split(" ", 1)[1]
        doc.add_heading(title, 2)
        doc.add_paragraph(desc)
        code(doc, rules)
        figure(doc, img, "Figure: %s architecture." % name, width=14.5)
        doc.add_paragraph(note)

        doc.add_heading("Python code", 3)
        doc.add_paragraph("The complete class from code/agents.py:")
        code(doc, inspect.getsource(cls), size=7.5)

        doc.add_heading("Reasoning in each situation", 3)
        figure(doc, "trace_%s.png" % key,
               "Figure: %s agent, reasoning trace in the three situations (percept, every reasoning step, "
               "decision, outcome)." % name, width=16.5)
        if key == "utility_based":
            figure(doc, "chart_utility.png", "Figure: utility of each route in each test. The agent rides the tallest bar.", width=16.5)
        if key == "learning":
            figure(doc, "chart_learning.png", "Figure: minutes the Learning agent learned for each route in each context.", width=16.5)

    # 5 -------------------------------------------------------------------
    doc.add_heading("5. Implementation", 1)
    table(doc, ["File", "Purpose"], [
        ["code/environment.py", "Road map, routes, test situations, percept, trip-time model, training days"],
        ["code/agents.py", "The seven agent classes; each has decide(percept) -> (route, reasoning)"],
        ["code/main.py", "Runs every agent on every test, prints the tables, writes results/results.txt and .json"],
        ["simulation/", "Interactive 3D simulation (Three.js): index.html, model.js, world.js, app.js"],
        ["tools/check_parity.js", "Checks that simulation/model.js makes exactly the same decisions, times and reasoning as the Python code"],
        ["tools/export_diagrams.py", "Draws the architecture diagrams"],
        ["tools/build_report.py", "Builds this report"],
    ], widths=[4.5, 12.1])
    doc.add_paragraph("Run the program with Python 3.8 or newer (no extra packages):")
    code(doc, "cd code\npython main.py            # result tables\npython main.py --details  # tables + every agent's reasoning")

    doc.add_heading("5.1 The 3D simulation", 2)
    doc.add_paragraph(
        "The simulation is a 3D model of the trip. A rider in a pink jacket on a motorbike with a delivery box "
        "leaves Cheezious and rides the chosen route to the UCP gate. The traffic keeps left, as in Pakistan. "
        "The scene shows Khayaban-e-Jinnah as a dual carriageway with a signal chowk, the Canal Bank Road along "
        "the Lahore Canal with its underpass, the inner streets of Block G, cars, rickshaws, buses and bikes, "
        "and lighting for the time of day. Test 2 has an evening sky and a traffic queue at the chowk. Test 3 has "
        "rain, a fallen tree across Khayaban-e-Jinnah and a flooded underpass with a stalled car.")
    bullets(doc, [
        "The decisions come from simulation/model.js, a line-by-line copy of the Python agents (verified by tools/check_parity.js).",
        "The rider's speed follows the model: free-flow speed on open road, crawling through the jam, pushing the bike through the flood. The total time equals the model's delivery time.",
        "The side panel shows the percept, the hidden truth (underpass flooded or dry) and the agent's reasoning line by line.",
        "Race mode sends all seven agents out at once. A custom mode lets you set the time, rain, jams, closures and flooding yourself.",
        "One simulated minute takes two seconds at 1x speed. Speeds of 2x, 4x and 8x and chase, map and orbit cameras are available.",
    ])
    figure(doc, "sim_test1_utility.png", "Figure: Test 1, the Utility-Based agent rides route A.", width=15.5)
    figure(doc, "sim_test2_goal_jam.png", "Figure: Test 2, the Goal-Based agent crawls through the jam at Jinnah chowk.", width=15.5)
    figure(doc, "sim_test3_start.png", "Figure: Test 3, rain, and Khayaban-e-Jinnah closed by a fallen tree.", width=15.5)
    figure(doc, "sim_test3_reflex_flood.png", "Figure: Test 3, the Simple Reflex agent is stuck in the flooded underpass.", width=15.5)
    figure(doc, "sim_test3_race.png", "Figure: Test 3 race. Five agents take route C; Simple Reflex and Goal-Based take B.", width=15.5)

    # 6 -------------------------------------------------------------------
    doc.add_heading("6. Results", 1)
    rows = []
    for a in res["agents"]:
        cells = ["%s %.1f%s" % (t["route"], t["minutes"], "" if t["on_time"] else " (late)") for t in a["tests"]]
        rows.append([a["agent"]] + cells + ["%d/3" % a["on_time"], "%.1f" % a["total_minutes"]])
    table(doc, ["Agent", "Test 1", "Test 2", "Test 3", "On time", "Total min"], rows,
          widths=[3.8, 2.4, 2.4, 2.8, 1.8, 2.0], highlight=3)
    doc.add_paragraph("Cell = chosen route and real delivery time in minutes. Deadline: 20 minutes.")
    figure(doc, "cmp_matrix.png", "Figure: the route each agent chose in each situation (colour = route).", width=15.5)

    best = {1: ("A", 10.0), 2: ("B", 10.4), 3: ("C", 14.4)}
    # the reasoning line that best explains each agent's choice
    DECIDING_STEP = {"simple_reflex": 0, "model_based": -1, "goal_based": 1, "utility_based": -1,
                     "learning": -1, "hierarchical": 1, "contextual": 1}
    for k, sc in enumerate(SCENARIOS):
        doc.add_heading("6.%d Test %d: %s" % (k + 1, sc["id"], sc["title"]), 2)
        doc.add_paragraph(sc["story"] + " Best possible: route %s, %.1f min." % best[sc["id"]])
        rows = []
        for a in res["agents"]:
            t = a["tests"][k]
            margin = DEADLINE_MIN - t["minutes"]
            rows.append([a["agent"], t["route"], "%.1f" % t["minutes"],
                         "On time" if t["on_time"] else "LATE",
                         ("%.1f spare" % margin) if margin >= 0 else ("%.1f late" % -margin),
                         "%+.1f" % (t["minutes"] - best[sc["id"]][1]),
                         t["reasoning"][DECIDING_STEP[a["key"]]]])
        table(doc, ["Agent", "Route", "Min", "Result", "Margin", "vs best", "Deciding step"], rows,
              widths=[2.7, 1.4, 1.2, 1.5, 2.1, 1.5, 6.2], highlight=3)

    doc.add_heading("6.4 Summary of the three situations", 2)
    bullets(doc, [
        "Test 1: every agent takes A (10.0 min). With nothing unusual, all architectures are equal.",
        "Test 2: the jam on A is visible on the app. Every agent except Goal-Based switches to B (10.4 min). "
        "Goal-Based still takes the shortest road A and needs 19.2 min. It is on time, but only just.",
        "Test 3: A is closed and the underpass on B is flooded, which cannot be seen. Simple Reflex (no memory) and "
        "Goal-Based (traffic and weather are not part of its goal) take B and arrive after 24.5 min, which is late. "
        "Every agent that keeps or learns a model of the rain (Model-Based, Utility, Learning, Hierarchical, Contextual) "
        "takes C and arrives in 14.4 min.",
    ])

    # 7 -------------------------------------------------------------------
    doc.add_heading("7. Comparison and the most suitable agent", 1)
    doc.add_heading("7.1 How the agents behaved differently", 2)
    doc.add_paragraph(
        "All seven agents received exactly the same percepts, so every difference below comes from the "
        "architecture alone. Test 1 does not separate them: with nothing unusual, every agent takes the shortest "
        "road A. Tests 2 and 3 show how each one reacts when something goes wrong.")
    figure(doc, "cmp_times.png", "Figure: delivery time per agent and situation; bar colour = route chosen, dashed line = 20-minute deadline.", width=16.5)
    bullets(doc, [
        "Simple Reflex reacts only to what it sees right now. It avoids the visible jam on A in Test 2, "
        "but in Test 3 its rule 'A closed, take B' sends the rider into the hidden flood (24.5 min, late).",
        "Model-Based Reflex uses the same kind of rules, but it remembers that it has been raining for 45 minutes. "
        "Its world model says the underpass floods after 30 minutes, so it picks C. Memory is the only difference from Simple Reflex, "
        "and it decides Test 3.",
        "Goal-Based plans a path to UCP on the map, but its cost is distance only. It ignores the jam in Test 2 "
        "(19.2 min, on time by only 0.8 min) and the flood in Test 3 (late). Reaching the goal is not the same as reaching it well.",
        "Utility-Based compares the expected outcome of every route. In Test 3 it gives B a utility of -42.5 "
        "(85% chance of a 24.5-minute flooded trip) and C a utility of 71.2, so it picks C.",
        "Learning makes the same choices as Utility-Based, but from experience: in its rain context, B averaged "
        "24.5 min and C 14.3 min over the training trips. It would need new training if the roads changed.",
        "Hierarchical splits the choice into layers. The strategic layer forbids congested and flood-prone roads, "
        "so the tactical layer is left with only B in Test 2 (C is slower) and only C in Test 3.",
        "Contextual recognises 'rush hour' and 'monsoon' and switches to the matching policy. It is right in "
        "all three tests, but only because these contexts were designed in advance.",
    ])
    figure(doc, "cmp_summary.png", "Figure: on-time deliveries and total delivery time per agent.", width=16)

    doc.add_heading("7.2 Feature comparison", 2)
    table(doc, ["Agent", "Handles hidden flood", "Handles traffic", "Needs training", "Design effort", "Explains its choice"], [
        ["Simple Reflex", "No", "Yes (visible jam)", "No", "Very low", "Rule that fired"],
        ["Model-Based Reflex", "Yes (rule)", "Yes", "No", "Low", "State + rule"],
        ["Goal-Based", "No", "No", "No", "Low", "Search path"],
        ["Utility-Based", "Yes (probability)", "Yes", "No", "Medium", "Utility of every route"],
        ["Learning", "Yes (learned)", "Yes", "Yes, 54 trips", "Medium", "Learned values"],
        ["Hierarchical", "Yes (constraint)", "Yes", "No", "Medium", "Each layer's decision"],
        ["Contextual", "Yes (context)", "Yes", "No", "Medium", "Context + policy"],
    ], widths=[3.2, 3.0, 2.6, 2.4, 2.3, 3.1], highlight=3)
    doc.add_paragraph(
        "Five agents deliver all three orders on time in 34.8 minutes in total, so the results alone do not decide.")
    doc.add_heading("7.3 The most suitable agent", 2)
    doc.add_paragraph(
        "We choose the Utility-Based agent as the most suitable for this problem, for these reasons:")
    bullets(doc, [
        "It deals with uncertainty directly. The flood cannot be observed, and the agent reasons with P(flood) "
        "instead of a yes/no rule. A rule that is wrong 15% of the time is still a rule; the utility agent knows how risky B is.",
        "It balances several goals at once: being on time, total minutes and rider safety. The weights can be "
        "changed (for example, a heavier penalty for lateness on a premium order) without rewriting the agent.",
        "It needs no training data, unlike the Learning agent, and does not depend on a fixed list of contexts, "
        "unlike the Contextual agent. A new mix of conditions (for example a jam on A and light rain together) is "
        "still scored correctly.",
        "Its reasoning is transparent: the utility of every route is shown, which makes each decision easy to explain to the rider or the dispatcher.",
    ])
    doc.add_paragraph(
        "The Learning agent is a strong second choice for a real company with a lot of trip data, and the "
        "Hierarchical design is a good way to add riding control on top of the Utility-Based route choice.")

    # 8 -------------------------------------------------------------------
    doc.add_heading("8. Conclusion", 1)
    doc.add_paragraph(
        "The simple architectures do well while the world is fully visible (Tests 1 and 2), but they fail when an "
        "important fact is hidden (Test 3). Agents that keep a model of the world, learn from experience or reason "
        "about outcomes avoid the flooded underpass and deliver on time. The Utility-Based agent combines this with "
        "explicit handling of uncertainty and adjustable goals, so it is the most suitable architecture for the "
        "Smart Food Delivery Agent. The 3D simulation makes these differences visible: the same order, the same "
        "roads, and seven riders that think in seven different ways.")

    os.makedirs(OUT_DIR, exist_ok=True)
    docx_path = os.path.join(OUT_DIR, NAME + ".docx")
    doc.save(docx_path)
    print("wrote", os.path.relpath(docx_path, ROOT))
    return docx_path


def to_pdf(docx_path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", OUT_DIR, docx_path],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    else:
        try:
            from docx2pdf import convert
        except ImportError:
            print("PDF skipped: install LibreOffice, or Microsoft Word + pip install docx2pdf")
            return
        convert(docx_path, os.path.join(OUT_DIR, NAME + ".pdf"))
    print("wrote", os.path.relpath(os.path.join(OUT_DIR, NAME + ".pdf"), ROOT))


if __name__ == "__main__":
    to_pdf(build())
